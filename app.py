import os, sqlite3, io
from datetime import datetime
from functools import wraps
from flask import Flask, request, redirect, url_for, session, render_template, send_file, flash
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
DB_PATH = os.environ.get("VACACIONES_DB", os.path.join(HERE, "vacaciones.db"))
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "cambiar-esta-clave")
app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "cambiar-esta-clave-secreta")

def connect():
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    return con

# Crea la tabla y, si ya existe, añade DNI sin borrar las respuestas anteriores.
with connect() as con:
    con.execute("""CREATE TABLE IF NOT EXISTS respuestas (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        nombre TEXT NOT NULL,
        vacaciones TEXT NOT NULL,
        inicio TEXT,
        fin TEXT,
        creado TEXT NOT NULL
    )""")
    columns = {row["name"] for row in con.execute("PRAGMA table_info(respuestas)").fetchall()}
    if "dni" not in columns:
        con.execute("ALTER TABLE respuestas ADD COLUMN dni TEXT NOT NULL DEFAULT ''")

def admin_only(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not session.get("admin"):
            return redirect(url_for("login"))
        return fn(*args, **kwargs)
    return wrapped

@app.route("/", methods=["GET", "POST"])
def home():
    if request.method == "POST":
        nombre = request.form.get("nombre", "").strip()
        dni = "".join(ch for ch in request.form.get("dni", "") if ch.isdigit())
        vacaciones = request.form.get("vacaciones", "")
        inicio = request.form.get("inicio", "").strip()
        fin = request.form.get("fin", "").strip()

        if not nombre or len(nombre) > 160:
            flash("Ingrese sus apellidos y nombres.")
            return redirect(url_for("home"))
        if len(dni) != 8:
            flash("Ingrese un DNI válido de 8 dígitos.")
            return redirect(url_for("home"))
        if vacaciones not in ("si", "no"):
            flash("Seleccione Sí o No.")
            return redirect(url_for("home"))

        with connect() as con:
            existing = con.execute("SELECT 1 FROM respuestas WHERE dni=? LIMIT 1", (dni,)).fetchone()
        if existing:
            flash("Este DNI ya tiene una respuesta registrada. Si necesita corregirla, comuníquese con el responsable.")
            return redirect(url_for("home"))

        if vacaciones == "si":
            try:
                d1 = datetime.strptime(inicio, "%Y-%m-%d").date()
                d2 = datetime.strptime(fin, "%Y-%m-%d").date()
                if d1.strftime("%Y-%m") != "2026-10" or d2.strftime("%Y-%m") != "2026-10" or d2 < d1:
                    raise ValueError
            except (ValueError, TypeError):
                flash("Indique fechas válidas de octubre de 2026. La fecha final no puede ser anterior a la inicial.")
                return redirect(url_for("home"))
            inicio_db, fin_db, vacaciones_db = inicio, fin, "Sí"
        else:
            inicio_db, fin_db, vacaciones_db = None, None, "No"

        with connect() as con:
            con.execute(
                "INSERT INTO respuestas (nombre, dni, vacaciones, inicio, fin, creado) VALUES (?, ?, ?, ?, ?, ?)",
                (nombre, dni, vacaciones_db, inicio_db, fin_db, datetime.now().isoformat(timespec="seconds"))
            )
        return render_template("gracias.html", nombre=nombre)
    return render_template("formulario.html")

@app.route("/admin/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        if request.form.get("password") == ADMIN_PASSWORD and ADMIN_PASSWORD != "cambiar-esta-clave":
            session["admin"] = True
            return redirect(url_for("admin"))
        flash("Contraseña incorrecta. Configure ADMIN_PASSWORD antes de usar el panel.")
    return render_template("login.html")

@app.route("/admin/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))

@app.route("/admin")
@admin_only
def admin():
    with connect() as con:
        rows = con.execute("SELECT * FROM respuestas ORDER BY id DESC").fetchall()
    return render_template("admin.html", rows=rows, total=len(rows),
        si=sum(row["vacaciones"] == "Sí" for row in rows),
        no=sum(row["vacaciones"] == "No" for row in rows))

@app.route("/admin/exportar.xlsx")
@admin_only
def exportar():
    with connect() as con:
        rows = con.execute(
            "SELECT id, nombre, dni, vacaciones, inicio, fin, creado FROM respuestas ORDER BY id"
        ).fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = "Vacaciones octubre 2026"
    ws.append(["N.º", "APELLIDOS Y NOMBRES", "DNI", "¿TUVO VACACIONES?", "FECHA DE INICIO", "FECHA DE FIN", "FECHA/HORA DE RESPUESTA"])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="155E75")
    for row in rows:
        ws.append([row["id"], row["nombre"], row["dni"], row["vacaciones"],
                   row["inicio"] or "", row["fin"] or "", row["creado"]])
    for i, width in enumerate([8, 36, 14, 22, 18, 18, 28], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, as_attachment=True, download_name="Respuestas_Vacaciones_Octubre_2026.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=False)
